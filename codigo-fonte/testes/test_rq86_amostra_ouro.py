"""RQ86: sorteio da amostra-ouro, planilhas de rotulagem e validacao (sem rede)."""

import shutil
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "validacao"))
import rq86_amostra_ouro as ouro


def releases_do_repo(nome, n=7):
    """``v0`` antes da janela, ``v1``..``vn`` na janela, mais um draft e uma pre-release."""
    linhas = [{"nome_completo": nome, "id": 1, "tag_name": "v0", "draft": False,
               "prerelease": False, "published_at": "2025-01-15T00:00:00Z",
               "dentro_janela": False}]
    for i in range(1, n + 1):
        linhas.append({"nome_completo": nome, "id": 10 + i, "tag_name": "v%d" % i,
                       "draft": False, "prerelease": False,
                       "published_at": "2026-%02d-10T00:00:00Z" % i, "dentro_janela": True})
    linhas += [
        {"nome_completo": nome, "id": 90, "tag_name": "rascunho", "draft": True,
         "prerelease": False, "published_at": "", "dentro_janela": False},
        {"nome_completo": nome, "id": 91, "tag_name": "v9-rc1", "draft": False,
         "prerelease": True, "published_at": "2026-08-01T00:00:00Z", "dentro_janela": True},
    ]
    return linhas


@pytest.fixture(scope="module")
def entrada(tmp_path_factory):
    """70 repositorios; ``o/novo`` so tem a 1a release da historia + 3 releases."""
    nomes = ["o/repo%02d" % i for i in range(69)] + ["o/novo"]
    metricas = pd.DataFrame({"nome_completo": nomes, "linguagem": "Go",
                             "estrelas": range(1000, 1070)})
    linhas = [l for nome in nomes[:-1] for l in releases_do_repo(nome)]
    linhas += [l for l in releases_do_repo("o/novo", n=4) if l["tag_name"] != "v0"]
    pasta = tmp_path_factory.mktemp("lab03_final")
    metricas.to_csv(pasta / "metricas_repositorios.csv", index=False)
    pd.DataFrame(linhas).to_csv(pasta / "releases.csv", index=False)
    return pasta


@pytest.fixture(scope="module")
def gerado_uma_vez(entrada, tmp_path_factory):
    saida = tmp_path_factory.mktemp("base") / "amostra_ouro"
    return saida, ouro.gerar(entrada, saida)


@pytest.fixture
def gerado(gerado_uma_vez, tmp_path):
    """Copia da saida gerada, para cada teste poder preencher e alterar a sua."""
    base, resumo = gerado_uma_vez
    saida = tmp_path / "amostra_ouro"
    shutil.copytree(base, saida)
    return saida, resumo


def preencher(planilha, tipo="biblioteca/framework", entregas="sim", corretiva="não",
              excecoes=None):
    """Preenche as colunas de rotulo; ``excecoes`` = {(aba, linha): {coluna: valor}}."""
    livro = load_workbook(planilha)
    for aba, valores in (("repositorios", {"tipo_projeto": tipo, "entregas_reais": entregas}),
                         ("releases", {"corretiva": corretiva})):
        planilha_aba = livro[aba]
        cabecalho = [c.value for c in planilha_aba[1]]
        for linha in range(2, planilha_aba.max_row + 1):
            for coluna, valor in valores.items():
                valor = (excecoes or {}).get((aba, linha), {}).get(coluna, valor)
                planilha_aba.cell(row=linha, column=cabecalho.index(coluna) + 1, value=valor)
    livro.save(planilha)


# ------------------------------------------------------------------ rotulos
@pytest.mark.parametrize("texto, esperado", [
    ("não", "nao"), ("Não ", "nao"), ("NAO", "nao"), ("nao", "nao"), ("sim", "sim"),
    ("incerto", "incerto"), ("talvez", None), ("", None), (None, None), (float("nan"), None),
])
def test_codigo_do_rotulo_de_entregas(texto, esperado):
    assert ouro.codigo_do_rotulo(texto, ouro.ENTREGAS_REAIS) == esperado


@pytest.mark.parametrize("texto, esperado", [
    ("aplicação/serviço", "aplicacao_servico"), ("Aplicacao/Servico", "aplicacao_servico"),
    ("aplicacao_servico", "aplicacao_servico"), ("ferramenta CLI", "ferramenta_cli"),
    ("biblioteca/framework", "biblioteca_framework"), ("outro", "outro"), ("lib", None),
])
def test_codigo_do_rotulo_de_tipo(texto, esperado):
    assert ouro.codigo_do_rotulo(texto, ouro.TIPOS_PROJETO) == esperado


# ------------------------------------------------------------------ sorteio
def test_sorteio_de_repositorios_e_o_df_sample_com_semente_42(entrada):
    metricas = pd.read_csv(entrada / "metricas_repositorios.csv")
    sorteados = ouro.sortear_repositorios(metricas)
    esperado = (metricas.sort_values("nome_completo").reset_index(drop=True)
                .sample(n=60, random_state=42)["nome_completo"].tolist())
    assert sorteados["nome_completo"].tolist() == esperado
    assert sorteados["id_repositorio"].tolist()[:3] == ["R01", "R02", "R03"]
    assert sorteados["nome_completo"].is_unique


def test_sorteio_nao_depende_da_ordem_do_csv(entrada):
    metricas = pd.read_csv(entrada / "metricas_repositorios.csv")
    embaralhada = metricas.sample(frac=1, random_state=7)
    assert (ouro.sortear_repositorios(metricas)["nome_completo"].tolist()
            == ouro.sortear_repositorios(embaralhada)["nome_completo"].tolist())


def test_outra_semente_gera_outro_sorteio(entrada):
    metricas = pd.read_csv(entrada / "metricas_repositorios.csv")
    assert (ouro.sortear_repositorios(metricas, semente=1)["nome_completo"].tolist()
            != ouro.sortear_repositorios(metricas)["nome_completo"].tolist())


def test_sorteio_recusa_amostra_pequena_ou_duplicada():
    with pytest.raises(ValueError, match="precisa de 60"):
        ouro.sortear_repositorios(pd.DataFrame({"nome_completo": ["o/a", "o/b"]}))
    with pytest.raises(ValueError, match="duplicados"):
        ouro.sortear_repositorios(pd.DataFrame({"nome_completo": ["o/a"] * 70}))


def test_releases_sorteadas_seguem_a_definicao_principal(entrada):
    metricas = pd.read_csv(entrada / "metricas_repositorios.csv")
    releases = pd.read_csv(entrada / "releases.csv", keep_default_na=False)
    repositorios, escolhidas = ouro.montar_sorteio(metricas, releases)
    assert set(escolhidas["tag_name"]) <= {"v%d" % i for i in range(1, 8)}  # sem draft/rc/v0
    por_repo = escolhidas.groupby("nome_completo").size()
    assert (por_repo.drop("o/novo", errors="ignore") == 5).all()
    assert set(escolhidas["id_repositorio"]) == set(repositorios["id_repositorio"])
    v1 = escolhidas[escolhidas["tag_name"] == "v1"]
    assert (v1["release_anterior"] == "v0").all()  # anterior fora da janela vale


def test_primeira_release_da_historia_nao_e_sorteada():
    repositorios = pd.DataFrame({"id_repositorio": ["R01"], "nome_completo": ["o/novo"]})
    releases = pd.DataFrame([l for l in releases_do_repo("o/novo", n=4) if l["tag_name"] != "v0"])
    escolhidas = ouro.sortear_releases(repositorios, releases)
    assert escolhidas["tag_name"].tolist() == ["v2", "v3", "v4"]  # v1 nao tem anterior
    assert escolhidas["id_release"].tolist() == ["R01-1", "R01-2", "R01-3"]


def test_sorteio_de_releases_e_reprodutivel(entrada):
    metricas = pd.read_csv(entrada / "metricas_repositorios.csv")
    releases = pd.read_csv(entrada / "releases.csv", keep_default_na=False)
    _, primeira = ouro.montar_sorteio(metricas, releases)
    _, segunda = ouro.montar_sorteio(metricas, releases.sample(frac=1, random_state=3))
    pd.testing.assert_frame_equal(primeira, segunda)


def test_links_escapam_tags_de_monorepo():
    metricas = pd.DataFrame({"nome_completo": ["o/r%02d" % i for i in range(60)]})
    linhas = []
    for nome in metricas["nome_completo"]:
        for i, tag in enumerate(["@s/pkg@1.0.0", "@s/pkg@1.1.0", "v 2"]):
            linhas.append({"nome_completo": nome, "id": i, "tag_name": tag, "draft": "False",
                           "prerelease": "False", "dentro_janela": "True",
                           "published_at": "2026-0%d-01T00:00:00Z" % (i + 1)})
    _, escolhidas = ouro.montar_sorteio(metricas, pd.DataFrame(linhas))
    links = set(escolhidas["link_commits"]) | set(escolhidas["link_release"])
    assert "https://github.com/o/r00/compare/@s/pkg@1.0.0...@s/pkg@1.1.0" in links
    assert "https://github.com/o/r00/releases/tag/v%202" in links


# ---------------------------------------------------------------- planilhas
def test_gerar_grava_sorteio_e_tres_planilhas(gerado):
    saida, resumo = gerado
    assert resumo["repositorios"] == 60
    assert (saida / "sorteio_repositorios.csv").exists()
    assert (saida / "sorteio_releases.csv").exists()
    assert sorted(p.name for p in (saida / "rotulos").iterdir()) == sorted(
        "rotulos_%s.xlsx" % a for a in ouro.AVALIADORES)


def test_planilha_e_cega_tem_listas_suspensas_e_links(gerado):
    saida, _ = gerado
    livro = load_workbook(saida / "rotulos" / "rotulos_bacelete.xlsx")
    assert livro.sheetnames == ["instrucoes", "repositorios", "releases"]
    assert "bacelete" in livro["instrucoes"]["A3"].value
    releases = livro["releases"]
    cabecalho = [c.value for c in releases[1]]
    assert cabecalho == [*ouro.COLUNAS_RELEASES, "corretiva", "observacoes"]
    # Nada da heuristica: sem mensagens de commit, sem classificacao SemVer.
    assert not {"mensagem", "patch", "heuristica", "semver"} & {c.lower() for c in cabecalho}
    listas = [v.formula1 for v in releases.data_validations.dataValidation]
    assert listas == ['"sim,não"']
    tipos = [v.formula1 for v in livro["repositorios"].data_validations.dataValidation]
    assert '"biblioteca/framework,aplicação/serviço,ferramenta CLI,outro"' in tipos
    link = releases.cell(row=2, column=cabecalho.index("link_commits") + 1)
    assert link.hyperlink.target.startswith("https://github.com/")


def test_gerar_nao_sobrescreve_rotulos_existentes(gerado, entrada):
    saida, _ = gerado
    with pytest.raises(FileExistsError):
        ouro.gerar(entrada, saida)
    assert ouro.gerar(entrada, saida, sobrescrever=True)["repositorios"] == 60


# ---------------------------------------------------------------- validacao
def test_planilha_vazia_lista_todas_as_celulas(gerado):
    saida, resumo = gerado
    resultado = ouro.validar_planilha(saida / "rotulos" / "rotulos_bacelete.xlsx")
    assert resultado["completa"] is False
    assert len(resultado["problemas"]) == 60 * 2 + resumo["releases"]


def test_planilha_preenchida_e_completa(gerado):
    saida, _ = gerado
    planilha = saida / "rotulos" / "rotulos_bacelete.xlsx"
    preencher(planilha)
    assert ouro.validar_planilha(planilha) == {"planilha": str(planilha), "problemas": [],
                                               "completa": True}
    abas = ouro.ler_planilha(planilha)
    assert set(abas["repositorios"]["tipo_projeto"]) == {"biblioteca_framework"}
    assert set(abas["releases"]["corretiva"]) == {"nao"}


@pytest.mark.parametrize("ausente", [None, float("nan"), pd.NA])
def test_validacao_de_abas_reutilizadas_rejeita_ausentes(gerado, ausente):
    saida, _ = gerado
    planilha = saida / "rotulos" / "rotulos_BernardoApl.xlsx"
    preencher(planilha)
    abas = ouro.ler_planilha(planilha)
    abas["repositorios"].loc[0, "tipo_projeto"] = ausente
    abas["repositorios"].loc[0, "tipo_projeto_original"] = ""
    resultado = ouro.validar_planilha(planilha, abas=abas)
    assert not resultado["completa"]
    assert resultado["problemas"] == ["repositorios R01: tipo_projeto vazio"]


def test_valor_fora_da_lista_e_celula_vazia_sao_apontados(gerado):
    saida, _ = gerado
    planilha = saida / "rotulos" / "rotulos_bacelete.xlsx"
    preencher(planilha, excecoes={("repositorios", 2): {"tipo_projeto": "lib"},
                                  ("releases", 3): {"corretiva": None}})
    problemas = ouro.validar_planilha(planilha)["problemas"]
    assert len(problemas) == 2
    assert "repositorios R01: tipo_projeto = 'lib'" in problemas[0]
    assert problemas[1].endswith("corretiva vazio")


def test_linha_alterada_em_relacao_ao_sorteio_e_apontada(gerado):
    saida, _ = gerado
    planilha = saida / "rotulos" / "rotulos_bacelete.xlsx"
    preencher(planilha)
    livro = load_workbook(planilha)
    livro["releases"]["D2"] = "outra-tag"
    livro.save(planilha)
    problemas = ouro.validar_planilha(planilha)["problemas"]
    assert problemas == ["aba releases: linhas diferentes do sorteio (alteradas, removidas ou "
                         "reordenadas)"]


def test_sorteio_de_referencia_ausente(gerado):
    saida, _ = gerado
    (saida / "sorteio_releases.csv").unlink()
    planilha = saida / "rotulos" / "rotulos_bacelete.xlsx"
    preencher(planilha)
    assert any("ausente" in p for p in ouro.validar_planilha(planilha)["problemas"])


# ---------------------------------------------------------------------- CLI
def test_cli_sem_amostra_final_falha_com_instrucao(tmp_path, capsys):
    assert ouro.principal(["gerar", "--entrada", str(tmp_path / "nada")]) == 1
    assert "RQ82" in capsys.readouterr().err


def test_cli_gerar_e_validar(entrada, tmp_path, capsys):
    saida = tmp_path / "cli"
    assert ouro.principal(["gerar", "--entrada", str(entrada), "--saida", str(saida)]) == 0
    assert ouro.principal(["gerar", "--entrada", str(entrada), "--saida", str(saida)]) == 1
    planilha = saida / "rotulos" / "rotulos_bacelete.xlsx"
    assert ouro.principal(["validar", str(planilha)]) == 1
    preencher(planilha)
    assert ouro.principal(["validar", str(planilha)]) == 0
    assert "completa" in capsys.readouterr().out
