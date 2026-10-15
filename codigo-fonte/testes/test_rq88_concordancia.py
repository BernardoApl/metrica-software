"""RQ88: kappa de Fleiss e consenso da amostra-ouro (sem rede)."""

import json
import random
import shutil
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "validacao"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import rq86_amostra_ouro as ouro
import rq88_concordancia as rq88
from test_rq86_amostra_ouro import releases_do_repo

A, B, C = ouro.AVALIADORES


def fleiss_manual(matriz, categorias):
    """Formula de Fleiss (1971) escrita a mao, independente do statsmodels."""
    n_itens, n_aval = matriz.shape
    contagens = [[list(linha).count(c) for c in categorias] for linha in matriz.itertuples(index=False)]
    p_j = [sum(linha[j] for linha in contagens) / (n_itens * n_aval) for j in range(len(categorias))]
    p_i = [(sum(x * x for x in linha) - n_aval) / (n_aval * (n_aval - 1)) for linha in contagens]
    p_barra, p_e = sum(p_i) / n_itens, sum(p * p for p in p_j)
    return (p_barra - p_e) / (1 - p_e)


# --------------------------------------------------------------------- kappa
def test_kappa_reproduz_o_exemplo_classico_de_fleiss():
    """Exemplo da Wikipedia: 10 itens, 14 avaliadores, 5 categorias -> kappa = 0,210."""
    tabela = [[0, 0, 0, 0, 14], [0, 2, 6, 4, 2], [0, 0, 3, 5, 6], [0, 3, 9, 2, 0],
              [2, 2, 8, 1, 1], [7, 7, 0, 0, 0], [3, 2, 6, 3, 0], [2, 5, 3, 2, 2],
              [6, 5, 2, 1, 0], [0, 2, 2, 3, 7]]
    categorias = ["c1", "c2", "c3", "c4", "c5"]
    linhas = [[c for c, n in zip(categorias, linha) for _ in range(n)] for linha in tabela]
    matriz = pd.DataFrame(linhas)
    assert rq88.kappa_fleiss(matriz, categorias) == pytest.approx(0.20993, abs=1e-5)


def test_kappa_igual_a_formula_manual_em_matrizes_aleatorias():
    sorteio = random.Random(88)
    categorias = list(ouro.TIPOS_PROJETO)
    for _ in range(50):
        linhas = [[sorteio.choice(categorias) for _ in range(3)] for _ in range(60)]
        matriz = pd.DataFrame(linhas, columns=[A, B, C])
        assert rq88.kappa_fleiss(matriz, categorias) == pytest.approx(
            fleiss_manual(matriz, categorias))


def test_concordancia_perfeita_e_kappa_1():
    matriz = pd.DataFrame({A: ["sim", "nao"] * 5, B: ["sim", "nao"] * 5, C: ["sim", "nao"] * 5})
    assert rq88.kappa_fleiss(matriz, ouro.CORRETIVA) == pytest.approx(1.0)


def test_todos_na_mesma_categoria_deixa_o_kappa_indefinido():
    matriz = pd.DataFrame({A: ["nao"] * 5, B: ["nao"] * 5, C: ["nao"] * 5})
    assert rq88.kappa_fleiss(matriz, ouro.CORRETIVA) is None
    resumo = rq88.resumo_concordancia("corretiva", matriz, ouro.CORRETIVA)
    assert resumo["interpretacao"] is None
    assert resumo["concordancia_total"] == 1.0


def test_categoria_nao_usada_entra_na_tabela():
    # Ninguem usou "incerto": o kappa deve ser o mesmo da formula com as 3 categorias.
    matriz = pd.DataFrame({A: ["sim", "nao", "sim"], B: ["sim", "nao", "nao"],
                           C: ["sim", "sim", "sim"]})
    assert rq88.kappa_fleiss(matriz, ouro.ENTREGAS_REAIS) == pytest.approx(
        fleiss_manual(matriz, list(ouro.ENTREGAS_REAIS)))


@pytest.mark.parametrize("kappa, esperado", [
    (-0.1, "pobre"), (0.0, "leve"), (0.20, "leve"), (0.21, "razoavel"), (0.40, "razoavel"),
    (0.41, "moderada"), (0.60, "moderada"), (0.61, "substancial"), (0.80, "substancial"),
    (0.81, "quase perfeita"), (1.0, "quase perfeita"), (None, None),
])
def test_interpretacao_de_landis_e_koch(kappa, esperado):
    assert rq88.interpretar_kappa(kappa) == esperado


def test_resumo_conta_unanimes_2_1_e_1_1_1():
    matriz = pd.DataFrame({A: ["outro", "outro", "outro", "outro"],
                           B: ["outro", "outro", "ferramenta_cli", "ferramenta_cli"],
                           C: ["outro", "ferramenta_cli", "biblioteca_framework", "outro"]},
                          index=["R01", "R02", "R03", "R04"])
    resumo = rq88.resumo_concordancia("tipo_projeto", matriz, ouro.TIPOS_PROJETO)
    assert (resumo["itens_unanimes"], resumo["itens_2_1"], resumo["itens_1_1_1"]) == (1, 2, 1)
    assert resumo["concordancia_total"] == 0.25
    assert resumo["distribuicao"] == {"biblioteca_framework": 1, "aplicacao_servico": 0,
                                      "ferramenta_cli": 3, "outro": 8}


# ------------------------------------------------------------------ consenso
@pytest.mark.parametrize("rotulos, esperado", [
    (["sim", "sim", "sim"], ("sim", "unanime")),
    (["sim", "nao", "sim"], ("sim", "maioria")),
    (["nao", "nao", "sim"], ("nao", "maioria")),
    (["sim", "nao", "incerto"], (None, "empate_pendente")),
])
def test_consenso_do_item(rotulos, esperado):
    assert rq88.consenso_do_item(rotulos) == esperado


def matriz_de_tipos():
    return pd.DataFrame({A: ["outro", "outro", "outro"],
                         B: ["outro", "ferramenta_cli", "ferramenta_cli"],
                         C: ["outro", "outro", "biblioteca_framework"]},
                        index=["R01", "R02", "R03"])


def test_decisao_de_reuniao_resolve_empate_e_pode_rever_maioria():
    decisoes = pd.DataFrame([
        {"dimensao": "tipo_projeto", "id_item": "R03", "rotulo": "ferramenta CLI",
         "justificativa": "README descreve um CLI"},
        {"dimensao": "tipo_projeto", "id_item": "R02", "rotulo": "ferramenta_cli",
         "justificativa": "revisto: o pacote e so o binario"},
        {"dimensao": "corretiva", "id_item": "R01-1", "rotulo": "sim", "justificativa": "x"},
    ])
    consenso = rq88.aplicar_consenso(matriz_de_tipos(), "tipo_projeto", ouro.TIPOS_PROJETO,
                                     decisoes)
    assert consenso["consenso"].tolist() == ["outro", "ferramenta_cli", "ferramenta_cli"]
    assert consenso["origem"].tolist() == ["unanime", "reuniao", "reuniao"]
    assert consenso.loc["R03", "justificativa"] == "README descreve um CLI"


@pytest.mark.parametrize("decisao, mensagem", [
    ({"id_item": "R99", "rotulo": "outro", "justificativa": "x"}, "nao e um item"),
    ({"id_item": "R03", "rotulo": "lib", "justificativa": "x"}, "invalido"),
    ({"id_item": "R03", "rotulo": "outro", "justificativa": " "}, "sem justificativa"),
])
def test_decisao_invalida_e_recusada(decisao, mensagem):
    decisoes = pd.DataFrame([dict(decisao, dimensao="tipo_projeto")])
    with pytest.raises(ValueError, match=mensagem):
        rq88.aplicar_consenso(matriz_de_tipos(), "tipo_projeto", ouro.TIPOS_PROJETO, decisoes)


def test_ler_decisoes(tmp_path):
    assert rq88.ler_decisoes(tmp_path / "nao_existe.csv").empty
    arquivo = tmp_path / "decisoes.csv"
    arquivo.write_text("dimensao,id_item\ntipo_projeto,R01\n", encoding="utf-8")
    with pytest.raises(ValueError, match="sem as colunas"):
        rq88.ler_decisoes(arquivo)
    arquivo.write_text("dimensao,id_item,rotulo,justificativa\n"
                       "tipo_projeto,R01,outro,a\ntipo_projeto,R01,outro,b\n", encoding="utf-8")
    with pytest.raises(ValueError, match="repete"):
        rq88.ler_decisoes(arquivo)


# ------------------------------------------------------------ fluxo completo
@pytest.fixture(scope="module")
def amostra_gerada(tmp_path_factory):
    nomes = ["o/repo%02d" % i for i in range(60)]
    entrada = tmp_path_factory.mktemp("lab03_final")
    pd.DataFrame({"nome_completo": nomes}).to_csv(entrada / "metricas_repositorios.csv",
                                                  index=False)
    pd.DataFrame([l for n in nomes for l in releases_do_repo(n)]).to_csv(
        entrada / "releases.csv", index=False)
    saida = tmp_path_factory.mktemp("base") / "amostra_ouro"
    ouro.gerar(entrada, saida)
    for avaliador in ouro.AVALIADORES:
        preencher_avaliador(saida / "rotulos" / ("rotulos_%s.xlsx" % avaliador), avaliador)
    return saida


#: Rotulos de cada avaliador por item: R01 e R02 divergem; o resto e unanime.
ROTULOS = {
    A: {"tipo_projeto": {"R01": "outro", "R02": "ferramenta CLI"}, "corretiva": {"R01-1": "sim"}},
    B: {"tipo_projeto": {"R01": "aplicação/serviço", "R02": "ferramenta CLI"},
        "corretiva": {"R01-1": "não"}},
    C: {"tipo_projeto": {"R01": "biblioteca/framework", "R02": "outro"},
        "corretiva": {"R01-1": "sim"}},
}


def preencher_avaliador(planilha, avaliador):
    padrao = {"tipo_projeto": "biblioteca/framework", "entregas_reais": "sim",
              "corretiva": "não"}
    livro = load_workbook(planilha)
    for aba in ("repositorios", "releases"):
        folha = livro[aba]
        cabecalho = [c.value for c in folha[1]]
        for linha in range(2, folha.max_row + 1):
            item = folha.cell(row=linha, column=1).value
            for coluna in ("tipo_projeto", "entregas_reais", "corretiva"):
                if coluna in cabecalho:
                    valor = ROTULOS[avaliador].get(coluna, {}).get(item, padrao[coluna])
                    folha.cell(row=linha, column=cabecalho.index(coluna) + 1, value=valor)
    livro.save(planilha)


@pytest.fixture
def pasta(amostra_gerada, tmp_path):
    destino = tmp_path / "amostra_ouro"
    shutil.copytree(amostra_gerada, destino)
    return destino


def test_fluxo_com_empate_pendente(pasta):
    resumo = rq88.executar(pasta)
    tipo = next(r for r in resumo["concordancia"] if r["dimensao"] == "tipo_projeto")
    assert (tipo["itens"], tipo["itens_unanimes"], tipo["itens_2_1"], tipo["itens_1_1_1"]) == (
        60, 58, 1, 1)
    entregas = next(r for r in resumo["concordancia"] if r["dimensao"] == "entregas_reais")
    assert entregas["kappa_fleiss"] is None  # todos "sim": kappa indefinido
    assert resumo["empates_sem_decisao"] == 1
    assert resumo["consenso_completo"] is False
    pauta = pd.read_csv(pasta / "pendencias_reuniao.csv")
    assert pauta[["dimensao", "id_item", "obrigatoria"]].values.tolist() == [
        ["tipo_projeto", "R01", True], ["tipo_projeto", "R02", False],
        ["corretiva", "R01-1", False]]
    assert pauta.loc[1, "maioria"] == "ferramenta_cli"


def test_fluxo_completo_com_decisao_de_reuniao(pasta):
    (pasta / "decisoes_reuniao.csv").write_text(
        "dimensao,id_item,rotulo,justificativa\n"
        "tipo_projeto,R01,aplicação/serviço,README: app desktop\n", encoding="utf-8")
    resumo = rq88.executar(pasta)
    assert resumo["consenso_completo"] is True
    assert resumo["origem_do_consenso"]["tipo_projeto"] == {"unanime": 58, "maioria": 1,
                                                           "reuniao": 1}
    repos = pd.read_csv(pasta / "consenso_repositorios.csv", keep_default_na=False)
    r01 = repos.set_index("id_repositorio").loc["R01"]
    assert r01["tipo_projeto_consenso"] == "aplicacao_servico"
    assert r01["tipo_projeto_origem"] == "reuniao"
    assert r01["tipo_projeto_bacelete"] == "aplicacao_servico"
    assert len(repos) == 60
    releases = pd.read_csv(pasta / "consenso_releases.csv", keep_default_na=False)
    r011 = releases.set_index("id_release").loc["R01-1"]
    assert (r011["corretiva_consenso"], r011["corretiva_origem"]) == ("sim", "maioria")
    assert r011["tag_name"].startswith("v")
    salvo = json.loads((pasta / "concordancia_fleiss.json").read_text(encoding="utf-8"))
    assert salvo["protocolo"].startswith("3-0 unanime")
    fleiss = pd.read_csv(pasta / "concordancia_fleiss.csv")
    assert fleiss["dimensao"].tolist() == ["tipo_projeto", "entregas_reais", "corretiva"]


def test_kappa_usa_os_rotulos_independentes_e_nao_o_consenso(pasta):
    antes = rq88.executar(pasta)["concordancia"]
    (pasta / "decisoes_reuniao.csv").write_text(
        "dimensao,id_item,rotulo,justificativa\ntipo_projeto,R01,outro,consenso\n",
        encoding="utf-8")
    assert rq88.executar(pasta)["concordancia"] == antes


def test_planilha_incompleta_impede_o_calculo(pasta):
    planilha = pasta / "rotulos" / ("rotulos_%s.xlsx" % B)
    livro = load_workbook(planilha)
    livro["releases"]["I2"] = None  # corretiva da primeira release
    livro.save(planilha)
    with pytest.raises(rq88.PlanilhaIncompleta, match=B):
        rq88.executar(pasta)
    (pasta / "rotulos" / ("rotulos_%s.xlsx" % C)).unlink()
    with pytest.raises(rq88.PlanilhaIncompleta, match="ausente"):
        rq88.executar(pasta)


def test_cli_codigos_de_saida(pasta, capsys):
    assert rq88.principal(["--pasta", str(pasta)]) == 2
    saida = capsys.readouterr().out
    assert "tipo_projeto" in saida and "1-1-1: 1" in saida and "indefinido" in saida
    (pasta / "decisoes_reuniao.csv").write_text(
        "dimensao,id_item,rotulo,justificativa\ntipo_projeto,R01,outro,x\n", encoding="utf-8")
    assert rq88.principal(["--pasta", str(pasta)]) == 0
    (pasta / "decisoes_reuniao.csv").write_text(
        "dimensao,id_item,rotulo,justificativa\ntipo_projeto,R01,lib,x\n", encoding="utf-8")
    assert rq88.principal(["--pasta", str(pasta)]) == 1
    (pasta / "rotulos" / ("rotulos_%s.xlsx" % A)).unlink()
    assert rq88.principal(["--pasta", str(pasta)]) == 1
    assert "incompletas" in capsys.readouterr().err
