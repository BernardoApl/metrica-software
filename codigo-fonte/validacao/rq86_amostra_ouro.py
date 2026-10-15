"""RQ86: sorteio da amostra-ouro e planilhas de rotulagem manual (secao 6 do enunciado).

    python codigo-fonte/validacao/rq86_amostra_ouro.py gerar
    python codigo-fonte/validacao/rq86_amostra_ouro.py validar dados/lab03_final/amostra_ouro/rotulos/rotulos_bacelete.xlsx

``gerar`` faz o sorteio **uma unica vez** para o grupo inteiro: os tres
avaliadores rotulam exatamente os mesmos 60 repositorios e as mesmas 300
releases, cada um na sua planilha. Com a mesma entrada e a mesma semente, o
sorteio sai identico, entao pode ser refeito e conferido por qualquer pessoa.

- 60 repositorios: ``df.sample(n=60, random_state=42)`` sobre a amostra final
  ordenada por nome (a ordem do CSV nao altera o sorteio);
- 5 releases por repositorio, com a mesma semente, entre as releases da
  definicao principal de deploy (publicada, nao draft, nao pre-release, dentro
  da janela) que tem release anterior. A primeira release da historia fica de
  fora: sem anterior, nao ha lista de commits para ler.

As planilhas sao **cegas**: trazem links para a release e para a lista de
commits (``compare`` com a release anterior), mas nenhuma saida da heuristica
de release corretiva (mudanca de SemVer, mensagens com ``fix``/``revert``).

``validar`` confere uma planilha preenchida antes do commit: valores fora das
listas, celulas vazias e linhas alteradas em relacao ao sorteio.
"""

from __future__ import annotations

import argparse
import sys
import unicodedata
from pathlib import Path
from typing import Optional
from urllib.parse import quote

import pandas as pd

RAIZ_CODIGO = Path(__file__).resolve().parents[1]
RAIZ_PROJETO = RAIZ_CODIGO.parent
sys.path.insert(0, str(RAIZ_CODIGO))
from metricas import lead_time  # noqa: E402

SEMENTE = 42
N_REPOSITORIOS = 60
N_RELEASES = 5
#: GitHub logins dos tres avaliadores (issues RQ85, RQ86 e RQ87).
AVALIADORES = ("BernardoApl", "bacelete", "Joaopedrotavaress")

ENTRADA_PADRAO = RAIZ_PROJETO / "dados" / "lab03_final"
SAIDA_PADRAO = ENTRADA_PADRAO / "amostra_ouro"

#: Codigo interno -> rotulo exibido na lista suspensa da planilha.
TIPOS_PROJETO = {
    "biblioteca_framework": "biblioteca/framework",
    "aplicacao_servico": "aplicação/serviço",
    "ferramenta_cli": "ferramenta CLI",
    "outro": "outro",
}
ENTREGAS_REAIS = {"sim": "sim", "nao": "não", "incerto": "incerto"}
CORRETIVA = {"sim": "sim", "nao": "não"}

#: Dimensoes rotuladas: (aba, coluna, valores permitidos).
DIMENSOES = (
    ("repositorios", "tipo_projeto", TIPOS_PROJETO),
    ("repositorios", "entregas_reais", ENTREGAS_REAIS),
    ("releases", "corretiva", CORRETIVA),
)
CHAVES = {"repositorios": ["id_repositorio", "nome_completo"],
          "releases": ["id_release", "nome_completo", "tag_name"]}

COLUNAS_REPOSITORIOS = ("id_repositorio", "nome_completo", "linguagem", "estrelas",
                        "link_repositorio", "link_releases")
COLUNAS_RELEASES = ("id_release", "id_repositorio", "nome_completo", "tag_name",
                    "release_anterior", "published_at", "link_release", "link_commits")

INSTRUCOES = """\
AMOSTRA-OURO — LAB03 (seção 6 do enunciado)

Avaliador(a): {avaliador}
Sorteio: {n_repos} repositórios e {n_releases} releases por repositório, semente {semente}.

REGRAS
1. Rotule SOZINHO(A). Não consulte os outros avaliadores, não compare planilhas e não
   discuta casos até os três terminarem e commitarem suas planilhas.
2. Não rode nem consulte a heurística automática de release corretiva (RQ89) antes de terminar.
3. Preencha só as colunas amarelas, usando as listas suspensas. Não altere as demais colunas.
4. Use "observacoes" para justificar casos difíceis; isso ajuda na reunião de consenso (RQ88).
5. Antes de commitar, rode:
   python codigo-fonte/validacao/rq86_amostra_ouro.py validar <esta planilha>

ABA "repositorios"
- tipo_projeto: o que o repositório é, principalmente.
    biblioteca/framework  -> código feito para ser importado por outros projetos (pacote, SDK, framework).
    aplicação/serviço     -> sistema que o usuário final executa ou acessa (app web/desktop/mobile, servidor, API).
    ferramenta CLI        -> programa usado pela linha de comando (dev tools, utilitários de terminal).
    outro                 -> nenhuma das anteriores (documentação, coleção de prompts/skills, dados, tutoriais...).
  Leia o README e a página do repositório. Se tiver mais de um papel, escolha o principal.
- entregas_reais: as releases deste repositório representam entregas reais ao usuário?
    sim      -> cada release disponibiliza uma versão nova que o usuário instala/usa (pacote publicado, binário, deploy).
    não      -> releases são marcos internos, automáticas sem conteúdo, snapshots, ou não chegam ao usuário.
    incerto  -> não dá para decidir com o que está disponível.

ABA "releases"
- corretiva: o objetivo PRINCIPAL desta release é corrigir um defeito recente?
    sim  -> a release existe sobretudo para consertar algo (bugfix/hotfix/revert de problema introduzido há pouco).
    não  -> traz funcionalidades, melhorias, atualizações de dependência, ou mistura em que correção não é o foco.
  Leia as release notes (link_release) e a lista de commits desde a release anterior (link_commits).
"""


# --------------------------------------------------------------------- utilitarios
def normalizar(valor) -> str:
    """Minusculas, sem acentos e sem espacos nas pontas ("Não " -> "nao")."""
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    texto = unicodedata.normalize("NFKD", str(valor)).encode("ascii", "ignore").decode()
    return texto.strip().casefold()


def codigo_do_rotulo(valor, permitidos: dict) -> Optional[str]:
    """Converte o texto da planilha no codigo interno; ``None`` se vazio ou invalido."""
    alvo = normalizar(valor)
    for codigo, rotulo in permitidos.items():
        if alvo in (normalizar(codigo), normalizar(rotulo)):
            return codigo
    return None


def _booleano(serie: pd.Series) -> pd.Series:
    return serie.astype(str).str.strip().str.lower().isin(("true", "1", "sim"))


def url_tag(tag: str) -> str:
    return quote(str(tag), safe="/@")


# -------------------------------------------------------------------------- sorteio
def sortear_repositorios(metricas: pd.DataFrame, n: int = N_REPOSITORIOS,
                         semente: int = SEMENTE) -> pd.DataFrame:
    """``df.sample(n, random_state=semente)`` sobre a amostra ordenada por nome."""
    if metricas["nome_completo"].duplicated().any():
        raise ValueError("A amostra final tem repositorios duplicados.")
    if len(metricas) < n:
        raise ValueError("A amostra tem %d repositorios; o sorteio precisa de %d."
                         % (len(metricas), n))
    ordenada = (metricas.assign(_chave=metricas["nome_completo"].str.casefold())
                .sort_values("_chave").drop(columns="_chave").reset_index(drop=True))
    sorteados = ordenada.sample(n=n, random_state=semente).reset_index(drop=True)
    sorteados.insert(0, "id_repositorio", ["R%02d" % (i + 1) for i in range(len(sorteados))])
    return sorteados


def releases_elegiveis(registros: list[dict], colunas) -> pd.DataFrame:
    """Deploys da janela (definicao principal) de um repositorio com release anterior conhecida."""
    linhas = [dict(release, release_anterior=anterior["tag_name"])
              for anterior, release in lead_time.pares_de_releases(registros)
              if anterior is not None]
    return pd.DataFrame(linhas, columns=[*colunas, "release_anterior"])


def sortear_releases(repositorios: pd.DataFrame, releases: pd.DataFrame,
                     k: int = N_RELEASES, semente: int = SEMENTE) -> pd.DataFrame:
    """``k`` releases por repositorio sorteado (todas, se houver menos de ``k`` elegiveis)."""
    releases = releases.copy()
    for coluna in ("draft", "prerelease", "dentro_janela"):
        releases[coluna] = _booleano(releases[coluna])
    releases["published_at"] = releases["published_at"].where(
        releases["published_at"].astype(str).str.strip() != "", None)
    por_repo: dict = {}
    for registro in releases.to_dict("records"):
        por_repo.setdefault(registro["nome_completo"], []).append(registro)
    blocos = []
    for _, repositorio in repositorios.iterrows():
        elegiveis = releases_elegiveis(por_repo.get(repositorio["nome_completo"], []),
                                       releases.columns)
        elegiveis = elegiveis.sort_values(["published_at", "id"]).reset_index(drop=True)
        escolhidas = elegiveis.sample(n=min(k, len(elegiveis)), random_state=semente)
        escolhidas = escolhidas.sort_values(["published_at", "id"])
        escolhidas.insert(0, "id_repositorio", repositorio["id_repositorio"])
        blocos.append(escolhidas)
    sorteio = pd.concat(blocos, ignore_index=True) if blocos else pd.DataFrame()
    sorteio.insert(0, "id_release", ["%s-%d" % (r, i + 1) for r, i in
                                     zip(sorteio["id_repositorio"],
                                         sorteio.groupby("id_repositorio").cumcount())])
    return sorteio


def montar_sorteio(metricas: pd.DataFrame, releases: pd.DataFrame,
                   semente: int = SEMENTE) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Tabelas de referencia do sorteio, com os links usados nas planilhas."""
    repositorios = sortear_repositorios(metricas, semente=semente)
    base = "https://github.com/"
    repositorios["link_repositorio"] = base + repositorios["nome_completo"]
    repositorios["link_releases"] = repositorios["link_repositorio"] + "/releases"
    for coluna in ("linguagem", "estrelas"):
        if coluna not in repositorios:
            repositorios[coluna] = ""
    escolhidas = sortear_releases(repositorios, releases, semente=semente)
    escolhidas["link_release"] = [
        "%s%s/releases/tag/%s" % (base, nome, url_tag(tag))
        for nome, tag in zip(escolhidas["nome_completo"], escolhidas["tag_name"])]
    escolhidas["link_commits"] = [
        "%s%s/compare/%s...%s" % (base, nome, url_tag(anterior), url_tag(tag))
        for nome, anterior, tag in zip(escolhidas["nome_completo"],
                                       escolhidas["release_anterior"], escolhidas["tag_name"])]
    return (repositorios[list(COLUNAS_REPOSITORIOS)].reset_index(drop=True),
            escolhidas[list(COLUNAS_RELEASES)].reset_index(drop=True))


# ------------------------------------------------------------------------ planilhas
def escrever_planilha(caminho: Path, avaliador: str, repositorios: pd.DataFrame,
                      releases: pd.DataFrame, semente: int = SEMENTE) -> None:
    """Planilha .xlsx com instrucoes, listas suspensas e links clicaveis."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    amarelo = PatternFill("solid", fgColor="FFF2CC")
    livro = Workbook()
    instrucoes = livro.active
    instrucoes.title = "instrucoes"
    texto = INSTRUCOES.format(avaliador=avaliador, n_repos=len(repositorios),
                              n_releases=N_RELEASES, semente=semente)
    for linha, conteudo in enumerate(texto.splitlines(), start=1):
        instrucoes.cell(row=linha, column=1, value=conteudo)
    instrucoes.column_dimensions["A"].width = 120
    instrucoes["A1"].font = Font(bold=True, size=13)

    abas = (
        ("repositorios", repositorios, (("tipo_projeto", TIPOS_PROJETO),
                                        ("entregas_reais", ENTREGAS_REAIS))),
        ("releases", releases, (("corretiva", CORRETIVA),)),
    )
    for nome_aba, dados, rotulos in abas:
        aba = livro.create_sheet(nome_aba)
        colunas = [*dados.columns, *(c for c, _ in rotulos), "observacoes"]
        aba.append(colunas)
        for registro in dados.itertuples(index=False):
            aba.append([None if pd.isna(v) else v for v in registro] + [None] * (len(rotulos) + 1))
        ultima = len(dados) + 1
        for indice, coluna in enumerate(colunas, start=1):
            letra = get_column_letter(indice)
            aba[letra + "1"].font = Font(bold=True)
            aba.column_dimensions[letra].width = 48 if coluna.startswith("link") else 22
            if coluna.startswith("link"):
                for linha in range(2, ultima + 1):
                    celula = aba["%s%d" % (letra, linha)]
                    if celula.value:
                        celula.hyperlink = celula.value
                        celula.font = Font(color="0563C1", underline="single")
        for coluna, permitidos in rotulos:
            letra = get_column_letter(colunas.index(coluna) + 1)
            validacao = DataValidation(type="list", allow_blank=True,
                                       formula1='"%s"' % ",".join(permitidos.values()),
                                       showErrorMessage=True, errorTitle="Valor inválido",
                                       error="Escolha um valor da lista.")
            aba.add_data_validation(validacao)
            validacao.add("%s2:%s%d" % (letra, letra, ultima))
            for linha in range(1, ultima + 1):
                aba["%s%d" % (letra, linha)].fill = amarelo
        letra_obs = get_column_letter(len(colunas))
        aba.column_dimensions[letra_obs].width = 50
        for linha in range(2, ultima + 1):
            aba["%s%d" % (letra_obs, linha)].alignment = Alignment(wrap_text=True)
        aba.freeze_panes = "C2"
    caminho.parent.mkdir(parents=True, exist_ok=True)
    livro.save(caminho)


def gerar(entrada: Path, saida: Path, semente: int = SEMENTE,
          avaliadores=AVALIADORES, sobrescrever: bool = False) -> dict:
    """Sorteia e grava ``sorteio_*.csv`` e uma planilha por avaliador."""
    metricas = pd.read_csv(entrada / "metricas_repositorios.csv")
    releases = pd.read_csv(entrada / "releases.csv", keep_default_na=False,
                           dtype={"tag_name": str})
    repositorios, escolhidas = montar_sorteio(metricas, releases, semente)
    planilhas = [saida / "rotulos" / ("rotulos_%s.xlsx" % a) for a in avaliadores]
    existentes = [p for p in planilhas if p.exists()]
    if existentes and not sobrescrever:
        raise FileExistsError("Planilhas ja existem e podem ter rotulos: %s. Use --sobrescrever "
                              "so se tiver certeza." % ", ".join(str(p) for p in existentes))
    saida.mkdir(parents=True, exist_ok=True)
    repositorios.to_csv(saida / "sorteio_repositorios.csv", index=False)
    escolhidas.to_csv(saida / "sorteio_releases.csv", index=False)
    for avaliador, planilha in zip(avaliadores, planilhas):
        escrever_planilha(planilha, avaliador, repositorios, escolhidas, semente)
    return {"repositorios": len(repositorios), "releases": len(escolhidas),
            "repositorios_com_menos_releases": int(
                (escolhidas.groupby("id_repositorio").size() < N_RELEASES).sum()),
            "planilhas": planilhas}


# ----------------------------------------------------------------------- leitura
def ler_planilha(caminho: Path) -> dict:
    """Le as abas ``repositorios`` e ``releases`` com os rotulos convertidos em codigos.

    Valores invalidos viram ``None`` em ``<dimensao>`` e o texto original fica
    em ``<dimensao>_original``, para o relatorio de validacao.
    """
    abas = pd.read_excel(caminho, sheet_name=["repositorios", "releases"], dtype=str)
    for aba, coluna, permitidos in DIMENSOES:
        tabela = abas[aba]
        if coluna not in tabela:
            raise ValueError("%s: aba %s sem a coluna %s." % (caminho, aba, coluna))
        tabela[coluna + "_original"] = tabela[coluna]
        tabela[coluna] = [codigo_do_rotulo(v, permitidos) for v in tabela[coluna]]
    return abas


def validar_planilha(caminho: Path, sorteio_dir: Optional[Path] = None,
                     abas: Optional[dict] = None) -> dict:
    """Problemas de uma planilha preenchida; ``completa`` se nao houver nenhum.

    ``abas`` evita reler o arquivo quando quem chama ja usou :func:`ler_planilha`.
    """
    abas = abas if abas is not None else ler_planilha(caminho)
    sorteio_dir = sorteio_dir or Path(caminho).resolve().parents[1]
    referencia = {"repositorios": sorteio_dir / "sorteio_repositorios.csv",
                  "releases": sorteio_dir / "sorteio_releases.csv"}
    problemas = []
    for aba, arquivo in referencia.items():
        if not arquivo.exists():
            problemas.append("sorteio de referencia ausente: %s" % arquivo)
            continue
        esperado = pd.read_csv(arquivo, dtype=str)[CHAVES[aba]].fillna("")
        obtido = abas[aba][CHAVES[aba]].fillna("")
        if not esperado.reset_index(drop=True).equals(obtido.reset_index(drop=True)):
            problemas.append("aba %s: linhas diferentes do sorteio (alteradas, removidas ou "
                             "reordenadas)" % aba)
    for aba, coluna, permitidos in DIMENSOES:
        tabela = abas[aba]
        chave = CHAVES[aba][0]
        for _, linha in tabela.iterrows():
            original = linha[coluna + "_original"]
            if linha[coluna] is not None:
                continue
            if normalizar(original) == "":
                problemas.append("%s %s: %s vazio" % (aba, linha[chave], coluna))
            else:
                problemas.append("%s %s: %s = %r nao esta em %s" % (
                    aba, linha[chave], coluna, original, list(permitidos.values())))
    return {"planilha": str(caminho), "problemas": problemas, "completa": not problemas}


# ---------------------------------------------------------------------------- CLI
def principal(argv=None) -> int:
    parser = argparse.ArgumentParser(description="RQ86: amostra-ouro e planilhas de rotulagem.")
    sub = parser.add_subparsers(dest="comando", required=True)
    p_gerar = sub.add_parser("gerar", help="sorteia 60 repositorios e gera as planilhas")
    p_gerar.add_argument("--entrada", type=Path, default=ENTRADA_PADRAO,
                         help="pasta com metricas_repositorios.csv e releases.csv (RQ82)")
    p_gerar.add_argument("--saida", type=Path, default=SAIDA_PADRAO)
    p_gerar.add_argument("--semente", type=int, default=SEMENTE)
    p_gerar.add_argument("--sobrescrever", action="store_true",
                         help="regrava planilhas existentes (APAGA rotulos ja feitos)")
    p_validar = sub.add_parser("validar", help="confere uma planilha preenchida")
    p_validar.add_argument("planilha", type=Path)
    args = parser.parse_args(argv)

    if args.comando == "gerar":
        if not (args.entrada / "metricas_repositorios.csv").exists():
            print("Erro: %s nao existe. Gere a amostra final antes (RQ82):\n"
                  "  python codigo-fonte/rq82_amostra_final.py --config config/rq82_amostra_final.json"
                  % (args.entrada / "metricas_repositorios.csv"), file=sys.stderr)
            return 1
        try:
            resumo = gerar(args.entrada, args.saida, args.semente, sobrescrever=args.sobrescrever)
        except (FileExistsError, ValueError) as erro:
            print("Erro: %s" % erro, file=sys.stderr)
            return 1
        print("Sorteio: %d repositorios, %d releases (semente %d)."
              % (resumo["repositorios"], resumo["releases"], args.semente))
        if resumo["repositorios_com_menos_releases"]:
            print("Aviso: %d repositorios tem menos de %d releases elegiveis."
                  % (resumo["repositorios_com_menos_releases"], N_RELEASES))
        for planilha in resumo["planilhas"]:
            print("  %s" % planilha)
        return 0

    resultado = validar_planilha(args.planilha)
    for problema in resultado["problemas"]:
        print("- " + problema)
    print("%s: %s" % (args.planilha, "completa" if resultado["completa"]
                      else "%d problema(s)" % len(resultado["problemas"])))
    return 0 if resultado["completa"] else 1


if __name__ == "__main__":
    raise SystemExit(principal())
